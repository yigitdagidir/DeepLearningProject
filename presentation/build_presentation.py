"""
Builds the final presentation (.pptx).

Like the report, it pulls the real numbers/figures from outputs/ so the slides
stay consistent with the code.  Run the experiments first, then:

    python presentation/build_presentation.py
"""

import os
import sys
import json

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image

# import the report's figure helpers
sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "report"))
import figures  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "outputs")
PRES_DIR = os.path.dirname(os.path.abspath(__file__))
ASSETS = figures.ASSETS

# ---- palette (Midnight Executive) ----
NAVY = RGBColor(0x1E, 0x27, 0x61)
INK = RGBColor(0x23, 0x2A, 0x3D)
ICE = RGBColor(0xCA, 0xDC, 0xFC)
ACCENT = RGBColor(0xF2, 0x9F, 0x3D)   # warm amber accent
GREY = RGBColor(0x5A, 0x61, 0x70)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

HEAD_FONT = "Georgia"
BODY_FONT = "Calibri"


# --------------------------------------------------------------------------- #
def load_json(path, default=None):
    try:
        with open(path) as f:
            return json.load(f)
    except Exception:
        return default


fusion = load_json(os.path.join(OUT, "fusion", "metrics.json"), {})
cfg = load_json(os.path.join(OUT, "fusion", "config.json"), {})
ablation = load_json(os.path.join(OUT, "ablation", "ablation.json"), [])
cv = load_json(os.path.join(OUT, "cross_validation.json"), {})
abl = {r["modality"]: r for r in ablation} if ablation else {}


def pct(x):
    return f"{100 * x:.1f}%" if isinstance(x, (int, float)) else "—"


# --------------------------------------------------------------------------- #
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height


def add_slide(bg=WHITE):
    slide = prs.slides.add_slide(BLANK)
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = bg
    return slide


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    return tb, tf


def set_run(run, text, size, color, bold=False, italic=False, font=BODY_FONT):
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font


def title_marker(slide):
    """Small amber square motif to the left of the title (our repeated motif)."""
    sq = slide.shapes.add_shape(1, Inches(0.6), Inches(0.62), Inches(0.18), Inches(0.18))
    sq.fill.solid(); sq.fill.fore_color.rgb = ACCENT
    sq.line.fill.background()


def slide_title(slide, text, color=NAVY):
    title_marker(slide)
    tb, tf = textbox(slide, 0.95, 0.45, 11.8, 0.9)
    p = tf.paragraphs[0]
    set_run(p.add_run(), text, 30, color, bold=True, font=HEAD_FONT)
    return tb


def bullets(slide, x, y, w, h, items, size=17, color=INK, gap=8):
    tb, tf = textbox(slide, x, y, w, h)
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        # support (text, level) or plain text
        if isinstance(it, tuple):
            text, lvl = it
        else:
            text, lvl = it, 0
        run = p.add_run()
        bullet_char = "•  " if lvl == 0 else "–  "
        set_run(run, bullet_char + text, size if lvl == 0 else size - 2,
                color if lvl == 0 else GREY)
        p.level = lvl
    return tb


def img_size(path):
    w, h = Image.open(path).size
    return w, h


def place_image(slide, path, x, y, width):
    w_px, h_px = img_size(path)
    h_in = width * h_px / w_px
    slide.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(width))
    return h_in


def place_stack(slide, paths, center_x, top, maxw, gap=0.22):
    """
    Stack equation images vertically, centred on `center_x`.
    Images are placed at their NATIVE size (rendered at 200 dpi -> /200 = inches)
    so that the text in every equation is the same physical size; only equations
    wider than `maxw` are scaled down.
    """
    y = top
    for pth in paths:
        w_px, h_px = img_size(pth)
        w_in = min(maxw, w_px / 200.0)
        h_in = w_in * h_px / w_px
        slide.shapes.add_picture(pth, Inches(center_x - w_in / 2), Inches(y),
                                 width=Inches(w_in))
        y += h_in + gap
    return y - top


def caption(slide, x, y, w, text):
    tb, tf = textbox(slide, x, y, w, 0.4)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    set_run(p.add_run(), text, 11, GREY, italic=True)


# =========================================================================== #
#  SLIDE 1 — TITLE (dark)                                                     #
# =========================================================================== #
s = add_slide(NAVY)
# accent square
sq = s.shapes.add_shape(1, Inches(0.9), Inches(2.05), Inches(0.5), Inches(0.5))
sq.fill.solid(); sq.fill.fore_color.rgb = ACCENT; sq.line.fill.background()

tb, tf = textbox(s, 0.9, 2.6, 11.5, 2.2)
p = tf.paragraphs[0]
set_run(p.add_run(), "Multi-Modal Classification of Images and Text", 40, WHITE,
        bold=True, font=HEAD_FONT)
p2 = tf.add_paragraph(); p2.space_before = Pt(6)
set_run(p2.add_run(), "Fusing a CNN and an RNN for fashion-item classification",
        20, ICE, italic=True)

tb2, tf2 = textbox(s, 0.9, 5.6, 11.5, 1.4)
p = tf2.paragraphs[0]
set_run(p.add_run(), "Deep Learning Project  •  2024–2025", 15, ICE)
p = tf2.add_paragraph(); p.space_before = Pt(6)
set_run(p.add_run(), "Member 1  ·  Member 2  ·  Member 3", 15, WHITE, bold=True)

# =========================================================================== #
#  SLIDE 2 — PROBLEM & OBJECTIVES                                             #
# =========================================================================== #
s = add_slide()
slide_title(s, "Problem & Objectives")
bullets(s, 0.6, 1.7, 6.1, 4.8, [
    "Fashion items come with BOTH a photo and a text description.",
    "Using only one modality throws away useful information.",
    ("A blurry photo may still have a clear caption — and vice-versa.", 1),
    "Goal: one model that uses image + text to classify the item.",
    "Objectives from the brief:",
    ("Classify images with a CNN", 1),
    ("Encode captions with embeddings + an RNN", 1),
    ("Fuse both modalities for the best accuracy", 1),
    ("Describe the optimal model mathematically", 1),
], size=17)
h = place_image(s, os.path.join(ASSETS, "architecture.png"), 7.0, 2.4, 5.9)
caption(s, 7.0, 2.4 + h + 0.05, 5.9, "Our multi-modal pipeline at a glance")

# =========================================================================== #
#  SLIDE 3 — DATASET                                                          #
# =========================================================================== #
s = add_slide()
slide_title(s, "Dataset: Fashion-Gen")
bullets(s, 0.6, 1.7, 6.0, 4.5, [
    "Fashion-Gen: clothing photos paired with style/category text.",
    "The image gives the visual signal; the caption gives the textual signal.",
    "Category label (e.g. SHOES, DRESSES) is the classification target.",
    "Reproducible demo set with the SAME structure ships with the code:",
    ("6 categories, learnable shapes + templated captions", 1),
    ("Lets us run the full pipeline in ~2 min, no big download", 1),
    ("One modality is corrupted per sample → lets us measure fusion", 1),
], size=16)
# montage
meta = os.path.join(ROOT, "data", "demo", "demo_metadata.json")
if os.path.exists(meta):
    mp = figures.dataset_montage(meta)
    h = place_image(s, mp, 6.9, 2.5, 6.0)
    caption(s, 6.9, 2.5 + h + 0.05, 6.0, "Sample images from the demo set (one per class)")

# =========================================================================== #
#  SLIDE 4 — ARCHITECTURE                                                     #
# =========================================================================== #
s = add_slide()
slide_title(s, "Model Architecture")
aw = 8.6
h = place_image(s, os.path.join(ASSETS, "architecture.png"), (13.333 - aw) / 2, 1.8, aw)
caption(s, (13.333 - aw) / 2, 1.8 + h + 0.12, aw,
        "Figure 1 — image branch (CNN) + text branch (RNN) + concatenation fusion + softmax")
bullets(s, 1.4, 1.8 + h + 0.55, 10.5, 0.8, [
    "Two encoders produce one vector each → concatenated → FC layers → softmax over classes.",
], size=16)

# =========================================================================== #
#  SLIDE 5 — IMAGE BRANCH (CNN)                                               #
# =========================================================================== #
s = add_slide()
slide_title(s, "Image Branch — CNN")
bullets(s, 0.6, 1.7, 6.0, 4.8, [
    "Convolution = learnable filters slide over the image, detecting local patterns.",
    ("Weight sharing → few parameters, translation-equivariant", 1),
    "ReLU adds the non-linearity; max-pooling down-samples.",
    "Transfer learning: ResNet pre-trained on ImageNet.",
    ("Drop its 1000-class head, keep the backbone as a feature extractor", 1),
    ("Residual (skip) connections let very deep nets train", 1),
    "Output: a 256-dim image vector v_img.",
], size=16)
place_stack(s, [os.path.join(ASSETS, n + ".png") for n in
                ["conv", "relu", "pool", "resnet"]],
            center_x=9.9, top=1.9, maxw=6.0, gap=0.34)

# =========================================================================== #
#  SLIDE 6 — TEXT BRANCH (RNN)                                                #
# =========================================================================== #
s = add_slide()
slide_title(s, "Text Branch — Embeddings + RNN")
bullets(s, 0.6, 1.7, 5.7, 4.8, [
    "Each word → dense embedding vector (trainable, or GloVe).",
    "An RNN reads the words and keeps a hidden state.",
    "We use an LSTM: a memory cell with 3 gates.",
    ("forget = what to keep, input = what to add, output = what to read", 1),
    ("Additive cell update → no vanishing gradient → long context", 1),
    "Bidirectional → uses past AND future words.",
    "GRU is a lighter 2-gate alternative.",
], size=16)
place_stack(s, [os.path.join(ASSETS, "embed.png"), os.path.join(ASSETS, "lstm.png")],
            center_x=9.7, top=1.9, maxw=6.3, gap=0.3)

# =========================================================================== #
#  SLIDE 7 — FUSION                                                           #
# =========================================================================== #
s = add_slide()
slide_title(s, "Multi-Modal Fusion")
bullets(s, 0.6, 1.7, 6.0, 4.8, [
    "Concatenate the two vectors into one joint vector z.",
    "Pass z through fully-connected (dense) layers + ReLU + dropout.",
    "Each FC unit sees features from BOTH modalities at once.",
    ("→ it can learn cross-modal correlations", 1),
    ("e.g. fire when image looks like a shoe AND caption says 'shoes'", 1),
    "This is what lets fusion beat either single modality.",
], size=17)
place_stack(s, [os.path.join(ASSETS, "concat.png"), os.path.join(ASSETS, "fusionfc.png"),
                os.path.join(ASSETS, "softmax.png")],
            center_x=9.9, top=2.2, maxw=6.0, gap=0.45)

# =========================================================================== #
#  SLIDE 8 — LOSS & OPTIMISATION                                              #
# =========================================================================== #
s = add_slide()
slide_title(s, "Loss & Optimisation")
bullets(s, 0.6, 1.7, 5.6, 4.8, [
    "Softmax turns logits into class probabilities.",
    "Cross-entropy loss = −log(prob of the correct class).",
    ("Small when confidently correct, large when wrong", 1),
    "We minimise it with Adam (adaptive gradient descent).",
    ("Keeps running mean & variance of gradients", 1),
    ("Per-parameter step size → fast, robust", 1),
    "Learning-rate step decay for fine-tuning near the end.",
], size=16)
place_stack(s, [os.path.join(ASSETS, "celoss.png"), os.path.join(ASSETS, "adam.png")],
            center_x=9.6, top=2.1, maxw=6.5, gap=0.5)

# =========================================================================== #
#  SLIDE 9 — REGULARISATION & VALIDATION                                      #
# =========================================================================== #
s = add_slide()
slide_title(s, "Regularisation & Validation")
bullets(s, 0.6, 1.7, 6.0, 4.8, [
    "Dropout in the fusion layers — randomly zero units while training.",
    ("Stops the model relying on any single feature", 1),
    "Weight decay (L2) + light image augmentation.",
    "k-fold cross-validation to check generalisation.",
    (f"Our 5-fold CV: mean {pct(cv.get('mean'))}, std "
     f"{cv.get('std', 0) * 100:.1f} pts → very stable", 1),
    "Random search to tune LR, batch size, RNN type, dropout, FC layers.",
], size=16)
place_stack(s, [os.path.join(ASSETS, "dropout.png")], center_x=9.9, top=2.5, maxw=5.9)
# a stat callout
tb, tf = textbox(s, 7.2, 3.7, 5.4, 2.0)
p = tf.paragraphs[0]
set_run(p.add_run(), pct(cv.get("mean")), 60, NAVY, bold=True, font=HEAD_FONT)
p2 = tf.add_paragraph()
set_run(p2.add_run(), "5-fold cross-validation accuracy", 15, GREY)

# =========================================================================== #
#  SLIDE 10 — RESULTS: TRAINING                                               #
# =========================================================================== #
s = add_slide()
slide_title(s, "Results — Training & Confusion")
h1 = place_image(s, os.path.join(OUT, "fusion", "training_curves.png"), 0.6, 1.9, 7.2)
caption(s, 0.6, 1.9 + h1 + 0.05, 7.2, "Loss & accuracy per epoch (train vs val)")
h2 = place_image(s, os.path.join(OUT, "fusion", "confusion_matrix.png"), 8.4, 1.8, 4.4)
caption(s, 8.4, 1.8 + h2 + 0.05, 4.4, "Confusion matrix (test set)")
tb, tf = textbox(s, 0.6, 6.6, 12.0, 0.7)
p = tf.paragraphs[0]
set_run(p.add_run(),
        f"Fused model: best val {pct(fusion.get('best_val_acc'))}  |  "
        f"test accuracy {pct(fusion.get('test_acc'))}", 16, INK, bold=True)

# =========================================================================== #
#  SLIDE 11 — RESULTS: FUSION ABLATION (key result)                          #
# =========================================================================== #
s = add_slide()
slide_title(s, "Does Fusion Help?  (Ablation)")
place_image(s, os.path.join(OUT, "ablation", "ablation_bar.png"), 0.6, 1.9, 5.6)
# big stat callouts on the right
def stat(slide, x, y, value, label, color):
    tb, tf = textbox(slide, x, y, 6.4, 1.2)
    p = tf.paragraphs[0]
    r = p.add_run(); set_run(r, value, 40, color, bold=True, font=HEAD_FONT)
    r2 = p.add_run(); set_run(r2, "   " + label, 17, GREY)

stat(s, 6.6, 2.1, pct(abl.get("image", {}).get("test_acc")), "Image only (CNN)", GREY)
stat(s, 6.6, 3.2, pct(abl.get("text", {}).get("test_acc")), "Text only (RNN)", GREY)
stat(s, 6.6, 4.3, pct(abl.get("both", {}).get("test_acc")), "Fusion (CNN + RNN)", NAVY)
tb, tf = textbox(s, 6.6, 5.6, 6.4, 1.4)
p = tf.paragraphs[0]
set_run(p.add_run(),
        "Each modality alone is limited (one modality is corrupted per sample). "
        "The FC fusion layer recovers the information neither carries alone.",
        16, INK)

# =========================================================================== #
#  SLIDE 12 — CONCLUSION (dark)                                               #
# =========================================================================== #
s = add_slide(NAVY)
sq = s.shapes.add_shape(1, Inches(0.9), Inches(0.8), Inches(0.4), Inches(0.4))
sq.fill.solid(); sq.fill.fore_color.rgb = ACCENT; sq.line.fill.background()
tb, tf = textbox(s, 1.5, 0.7, 11.0, 1.0)
set_run(tf.paragraphs[0].add_run(), "Conclusion", 34, WHITE, bold=True, font=HEAD_FONT)

bullets(s, 1.0, 2.0, 11.5, 4.0, [
    "We built and mathematically described a CNN + RNN multi-modal classifier.",
    "Transfer learning (ResNet) + a gated RNN (LSTM/GRU) + concatenation fusion.",
    f"Fusion lifted test accuracy from {pct(abl.get('image', {}).get('test_acc'))} "
    f"(image) and {pct(abl.get('text', {}).get('test_acc'))} (text) to "
    f"{pct(abl.get('both', {}).get('test_acc'))} (fused).",
    "The fully-connected fusion layer is what captures the image–text correlations.",
    "Future work: attention-based fusion, full Fashion-Gen training.",
], size=18, color=ICE, gap=12)

tb, tf = textbox(s, 1.0, 6.6, 11.5, 0.7)
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
set_run(p.add_run(), "Thank you — Questions?", 20, ACCENT, bold=True, font=HEAD_FONT)

# =========================================================================== #
out_path = os.path.join(PRES_DIR, "Presentation_MultiModal_Classification.pptx")
prs.save(out_path)
print("saved presentation to", out_path)
